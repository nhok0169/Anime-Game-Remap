import os
from typing import Dict, List, Optional, Sequence, Tuple, Union

from .VGMatcher import VGMatch


# what one draft cell can hold for a target: an index (single-component target), or a
#   (component, index) pair (multi-component target), or nothing
DraftTarget = Union[None, int, Tuple[str, int]]


class ProposedSheet():
    """
    One remap direction as it goes into a draft: the vertex groups of one source component onto
    a target character (all of its components)

    Parameters
    ----------
    fromName: :class:`str`
        Column A's header --- the source character's name with its component appended
        (``YelanTranquilBody``), or just the name for a single-component source

    toName: :class:`str`
        The target character's name

    version: Optional[:class:`str`]
        The game version, for the sheet's title

    matches: List[:class:`VGMatch`]
        The source component's matches, in index order

    toComponents: Sequence[:class:`str`]
        The target's component names in draw order --- one target column each (``[""]`` for a
        single-component target, which gives the classic four-column sheet)

    Attributes
    ----------
    fromName: :class:`str`
        Column A's header

    toName: :class:`str`
        The target character's name

    version: Optional[:class:`str`]
        The game version

    matches: List[:class:`VGMatch`]
        The matches

    toComponents: List[:class:`str`]
        The target's component names
    """

    def __init__(self, fromName: str, toName: str, version: Optional[str], matches: List[VGMatch], toComponents: Sequence[str] = ("",)):
        self.fromName = fromName
        self.toName = toName
        self.version = version
        self.matches = matches
        self.toComponents = list(toComponents)

    def __iter__(self):
        # so the older (fromName, toName, version, matches) unpacking keeps working
        return iter((self.fromName, self.toName, self.version, self.matches))

    def rows(self) -> List[Tuple[int, List[Optional[int]], Optional[float], str]]:
        """
        The sheet's rows: ``(fromIndex, [toIndex per target component], uncertainty, comment)``

        Returns
        -------
        List[Tuple[:class:`int`, List[Optional[:class:`int`]], Optional[:class:`float`], :class:`str`]]
            The rows
        """

        result = []
        for match in self.matches:
            uncertainty: Optional[float] = round(match.uncertainty, 3)
            if (uncertainty == 0):
                uncertainty = None
            targets: List[Optional[int]] = [None] * len(self.toComponents)
            if (match.toGroup is not None and match.toComponent in self.toComponents):
                targets[self.toComponents.index(match.toComponent)] = match.toIndex
            result.append((match.fromIndex, targets, uncertainty, match.comment()))
        return result


class DraftWriter():
    """
    Writes proposed remaps into an Excel workbook in the format of the drafts under
    ``Data/RemapDrafts/`` (see that folder's ``README.md``): one sheet per remap direction, with
    the columns ``[from mod] | [to mod] | Uncertainty | Comments`` and duplicate ``to`` indices
    highlighted in yellow :raw-html:`<br />` :raw-html:`<br />`

    A target of several components gets **one column per component**, headed by the target's
    name with the component appended (``YelanTranquilBody | YelanTranquilBang | YelanTranquilEye``),
    exactly one of which is filled per row; a source of several components gets **one sheet per
    component**, headed the same way (``YelanTranquilBody`` in column A)

    Parameters
    ----------
    path: :class:`str`
        Where to write the workbook. If the file exists, sheets for the same remap directions are
        replaced and every other sheet is kept

    Attributes
    ----------
    path: :class:`str`
        Where the workbook is written
    """

    MaxSheetTitleLen = 31
    DuplicateFill = "FFFF00"

    ColumnWidths = {"A": 14, "B": 16, "C": 12, "D": 110}
    TargetColumnWidth = 18

    UncertaintyHeader = "Uncertainty"
    CommentsHeader = "Comments"

    # Every workbook this writer produces carries this sheet, so a proposal can be told apart from
    #   a hand-made draft (the benchmark refuses to score the tool against its own output)
    AboutSheet = "About"
    AboutText = "Proposed by Tools/VGRemapFinder. Remove this sheet once the proposal has been checked and is a draft in its own right."

    # A single proposed sheet added to a hand-made workbook is marked in the header cell after
    #   'Comments' instead, so the benchmark can skip the sheet without disowning the workbook
    ProposedMarker = "Proposed by Tools/VGRemapFinder"

    def __init__(self, path: str):
        self.path = path

    @classmethod
    def sheetTitle(cls, fromName: str, toName: str, version: Optional[str] = None) -> str:
        """
        Names a sheet the way the hand-made drafts do (``"V4.4 - Ganyu to GanyuTwilight"``),
        trimmed to Excel's 31 character limit

        Parameters
        ----------
        fromName: :class:`str`
            The name of the mod being remapped

        toName: :class:`str`
            The name of the remapped mod

        version: Optional[:class:`str`]
            The game version the remap targets

        Returns
        -------
        :class:`str`
            The sheet's title
        """

        body = f"{fromName} to {toName}"
        title = body if (version is None) else f"V{version} - {body}"

        if (len(title) > cls.MaxSheetTitleLen):
            title = title[:cls.MaxSheetTitleLen]
        return title

    def write(self, sheets: Sequence[Union[ProposedSheet, Tuple[str, str, Optional[str], List[VGMatch]]]]):
        """
        Writes the workbook

        Parameters
        ----------
        sheets: Sequence[Union[:class:`ProposedSheet`, Tuple[:class:`str`, :class:`str`, Optional[:class:`str`], List[:class:`VGMatch`]]]]
            The sheets to write (a plain ``(fromName, toName, version, matches)`` tuple is taken
            as a single-component target)
        """

        try:
            import openpyxl
        except ImportError as e:
            raise ImportError("Writing the draft workbook needs 'openpyxl' (pip install openpyxl)") from e

        if (os.path.isfile(self.path)):
            workbook = openpyxl.load_workbook(self.path)
        else:
            workbook = openpyxl.Workbook()
            workbook.remove(workbook.active)

        if (self.AboutSheet not in workbook.sheetnames):
            about = workbook.create_sheet(self.AboutSheet, 0)
            about.append([self.AboutText])
            about.column_dimensions["A"].width = 120

        for sheet in sheets:
            if (not isinstance(sheet, ProposedSheet)):
                sheet = ProposedSheet(*sheet)
            self.addSheet(workbook, sheet.fromName, sheet.toName, sheet.version, sheet.rows(), toComponents = sheet.toComponents)

        folder = os.path.dirname(os.path.abspath(self.path))
        os.makedirs(folder, exist_ok = True)
        workbook.save(self.path)

    @classmethod
    def addSheet(cls, workbook, fromName: str, toName: str, version: Optional[str],
                 rows: Sequence[Tuple[int, Union[None, int, Sequence[Optional[int]]], Optional[float], Optional[str]]],
                 marker: Optional[str] = None, title: Optional[str] = None, toComponents: Sequence[str] = ("",)):
        """
        Adds one remap-direction sheet, in the drafts' format, to an open ``openpyxl`` workbook,
        replacing a sheet of the same title

        Parameters
        ----------
        workbook: :class:`openpyxl.Workbook`
            The workbook to add to

        fromName: :class:`str`
            The name of the mod being remapped (column A's header)

        toName: :class:`str`
            The name of the remapped mod (the target columns' header, with each component appended)

        version: Optional[:class:`str`]
            The game version, for the sheet's title

        rows: Sequence[Tuple[:class:`int`, Union[None, :class:`int`, Sequence[Optional[:class:`int`]]], Optional[:class:`float`], Optional[:class:`str`]]]
            The rows, each ``(fromIndex, toIndex or [toIndex per target component], uncertainty, comment)``

        marker: Optional[:class:`str`]
            Text to put in the header cell after ``Comments``, marking the sheet as a proposal
            rather than a hand-made draft (``None`` for no mark)

        title: Optional[:class:`str`]
            The sheet's title. Defaults to :meth:`sheetTitle`

        toComponents: Sequence[:class:`str`]
            The target's component names, one column each (``[""]`` for a single-component target)

        Returns
        -------
        The sheet
        """

        from openpyxl.formatting.rule import Rule
        from openpyxl.styles import Font, PatternFill
        from openpyxl.styles.differential import DifferentialStyle
        from openpyxl.utils import get_column_letter

        toComponents = list(toComponents) or [""]

        if (title is None):
            title = cls.sheetTitle(fromName, toName, version)
        title = title[:cls.MaxSheetTitleLen]
        if (title in workbook.sheetnames):
            workbook.remove(workbook[title])

        sheet = workbook.create_sheet(title)
        header = [fromName] + [toName + component for component in toComponents] + [cls.UncertaintyHeader, cls.CommentsHeader]
        sheet.append(header)
        for cell in sheet[1]:
            cell.font = Font(bold = True)

        if (marker):
            sheet.cell(row = 1, column = len(header) + 1).value = marker

        for fromIndex, targets, uncertainty, comment in rows:
            if (targets is None or isinstance(targets, int)):
                targets = [targets] + [None] * (len(toComponents) - 1)
            sheet.append([fromIndex] + list(targets) + [uncertainty, comment])

        lastRow = len(rows) + 1
        if (lastRow >= 2):
            fill = PatternFill(start_color = cls.DuplicateFill, end_color = cls.DuplicateFill, fill_type = "solid")
            for column in range(2, 2 + len(toComponents)):
                letter = get_column_letter(column)
                rule = Rule(type = "duplicateValues", dxf = DifferentialStyle(fill = fill))
                sheet.conditional_formatting.add(f"{letter}2:{letter}{lastRow}", rule)

        sheet.column_dimensions["A"].width = cls.ColumnWidths["A"]
        for column in range(2, 2 + len(toComponents)):
            sheet.column_dimensions[get_column_letter(column)].width = cls.ColumnWidths["B"] if (len(toComponents) == 1) else cls.TargetColumnWidth
        sheet.column_dimensions[get_column_letter(2 + len(toComponents))].width = cls.ColumnWidths["C"]
        sheet.column_dimensions[get_column_letter(3 + len(toComponents))].width = cls.ColumnWidths["D"]

        sheet.freeze_panes = "A2"
        return sheet

    @classmethod
    def parseHeader(cls, header: Sequence, fromName: str, toName: str) -> Optional[Tuple[int, List[Tuple[int, str]]]]:
        """
        Whether a header row is the sheet for ``fromName -> toName``, and if so which of its
        columns are the target columns

        Parameters
        ----------
        header: Sequence
            The header row's values

        fromName: :class:`str`
            The expected column A header

        toName: :class:`str`
            The target's name, which every target column's header starts with

        Returns
        -------
        Optional[Tuple[:class:`int`, List[Tuple[:class:`int`, :class:`str`]]]]
            ``None`` if the header is some other sheet's; otherwise the 0-based column of
            ``Uncertainty`` (or where it would be) and the target columns as ``(column,
            component)``, the component being the header minus ``toName``
        """

        if (header is None or len(header) < 2 or header[0] is None or str(header[0]).strip() != fromName):
            return None

        targets: List[Tuple[int, str]] = []
        column = 1
        while (column < len(header)):
            value = header[column]
            text = "" if (value is None) else str(value).strip()
            if (text == cls.UncertaintyHeader or text == cls.CommentsHeader or text == ""):
                break
            if (not text.startswith(toName)):
                return None
            targets.append((column, text[len(toName):]))
            column += 1

        if (not targets):
            return None
        return column, targets

    @classmethod
    def isProposedSheet(cls, sheet) -> bool:
        """
        Whether a sheet was marked as a proposal by :meth:`addSheet`

        Parameters
        ----------
        sheet: :class:`openpyxl.worksheet.worksheet.Worksheet`
            The sheet (a read-only one works too)

        Returns
        -------
        :class:`bool`
            Whether the sheet's header row carries the :attr:`ProposedMarker`
        """

        try:
            header = next(sheet.iter_rows(min_row = 1, max_row = 1, values_only = True), None)
        except Exception:
            return False
        if (header is None):
            return False
        return any(value is not None and str(value).startswith(cls.ProposedMarker) for value in header)

    @classmethod
    def isProposal(cls, path: str) -> bool:
        """
        Whether a workbook was written by this tool (it carries the :attr:`AboutSheet`), as
        opposed to being a hand-made draft

        Parameters
        ----------
        path: :class:`str`
            The workbook to check

        Returns
        -------
        :class:`bool`
            Whether the workbook is one of this tool's proposals
        """

        try:
            import openpyxl
        except ImportError as e:
            raise ImportError("Reading a draft workbook needs 'openpyxl' (pip install openpyxl)") from e

        workbook = openpyxl.load_workbook(path, read_only = True)
        return cls.AboutSheet in workbook.sheetnames

    @classmethod
    def readSheet(cls, path: str, fromName: str, toName: str) -> Dict[int, DraftTarget]:
        """
        Reads one remap direction back out of a draft workbook, found by its header row

        Parameters
        ----------
        path: :class:`str`
            The workbook to read

        fromName: :class:`str`
            The name in the header of column A

        toName: :class:`str`
            The target's name: the header of column B for a single-component target, or the
            prefix every target column's header shares for a multi-component one

        Returns
        -------
        Dict[:class:`int`, :data:`DraftTarget`]
            The remap, source index to target: an index when the sheet has the classic single
            target column headed exactly ``toName``, a ``(component, index)`` pair otherwise,
            ``None`` where the draft left the row blank

        Raises
        ------
        :class:`KeyError`
            If no sheet has that header
        """

        try:
            import openpyxl
        except ImportError as e:
            raise ImportError("Reading a draft workbook needs 'openpyxl' (pip install openpyxl)") from e

        workbook = openpyxl.load_workbook(path, read_only = True, data_only = True)
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(min_row = 1, values_only = True)
            header = next(rows, None)
            parsed = cls.parseHeader(header, fromName, toName)
            if (parsed is None):
                continue

            _, targets = parsed
            classic = len(targets) == 1 and targets[0][1] == ""

            result: Dict[int, DraftTarget] = {}
            for row in rows:
                if (not row or row[0] is None or str(row[0]).strip() == ""):
                    continue
                try:
                    fromIndex = int(row[0])
                except (TypeError, ValueError):
                    continue

                value: DraftTarget = None
                for column, component in targets:
                    if (column < len(row) and row[column] is not None and str(row[column]).strip() != ""):
                        try:
                            index = int(row[column])
                        except (TypeError, ValueError):
                            continue
                        value = index if (classic) else (component, index)
                        break
                result[fromIndex] = value
            return result

        raise KeyError(f"'{path}' has no sheet whose header row is '{fromName}' | '{toName}...'")
